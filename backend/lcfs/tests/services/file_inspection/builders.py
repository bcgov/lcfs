"""Build small, real Office files in memory for upload tests."""

import io
import zipfile
from pathlib import Path

import docx
import xlwt
from openpyxl import Workbook

FIXTURES = Path(__file__).with_name("fixtures")


def xlsx_bytes() -> bytes:
    buffer = io.BytesIO()
    workbook = Workbook()
    workbook.active["A1"] = "Fuel"
    workbook.save(buffer)
    return buffer.getvalue()


def docx_bytes(text: str = "hello") -> bytes:
    buffer = io.BytesIO()
    document = docx.Document()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def xls_bytes() -> bytes:
    buffer = io.BytesIO()
    workbook = xlwt.Workbook()
    workbook.add_sheet("Sheet1").write(0, 0, "Fuel")
    workbook.save(buffer)
    return buffer.getvalue()


def rewrite_zip(data: bytes, edits: dict) -> bytes:
    """
    Copy a zip, passing each named entry's bytes through its edit function.
    Names not already in the zip are added with edit(b"").
    """
    edits = dict(edits)
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info)
            if info.filename in edits:
                content = edits.pop(info.filename)(content)
            target.writestr(info.filename, content)
        for name, edit in edits.items():
            target.writestr(name, edit(b""))
    return out.getvalue()


def sample_vba_project() -> bytes:
    return (FIXTURES / "xlsxwriter_vbaProject.bin").read_bytes()


def xlsm_bytes(vba_project: bytes | None = None, extra: dict | None = None) -> bytes:
    """A macro-enabled workbook; with vba_project=None it carries no macros."""
    edits = {
        "[Content_Types].xml": lambda c: c.replace(
            b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
            b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
        ).replace(
            b"</Types>",
            b'<Default Extension="bin" '
            b'ContentType="application/vnd.ms-office.vbaProject"/></Types>',
        )
    }
    if vba_project is not None:
        edits["xl/_rels/workbook.xml.rels"] = lambda c: c.replace(
            b"</Relationships>",
            b'<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/'
            b'2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>',
        )
        edits["xl/vbaProject.bin"] = lambda _: vba_project
    edits.update(extra or {})
    return rewrite_zip(xlsx_bytes(), edits)

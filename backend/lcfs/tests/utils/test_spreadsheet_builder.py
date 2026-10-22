from datetime import date
from io import BytesIO

import numpy as np
import xlrd
from openpyxl import load_workbook

from lcfs.utils.spreadsheet_builder import SpreadsheetBuilder, SpreadsheetColumn


def test_xls_writer_formats_native_date_cells():
    builder = SpreadsheetBuilder(file_format="xls")
    builder.add_sheet(
        sheet_name="Dates",
        columns=[
            SpreadsheetColumn("Transaction date", "date"),
            SpreadsheetColumn("Description", "text"),
        ],
        rows=[[date(2026, 9, 17), "Transfer"]],
    )

    workbook = xlrd.open_workbook(
        file_contents=builder.build_spreadsheet(),
        formatting_info=True,
    )
    sheet = workbook.sheet_by_name("Dates")

    date_cell = sheet.cell(1, 0)
    xf = workbook.xf_list[date_cell.xf_index]
    format_key = xf.format_key

    assert date_cell.ctype == xlrd.XL_CELL_DATE
    assert workbook.format_map[format_key].format_str.lower() == "yyyy-mm-dd"
    assert xlrd.xldate_as_datetime(date_cell.value, workbook.datemode).date() == date(
        2026, 9, 17
    )


def test_numeric_numpy_scalars_and_totals_remain_numeric_in_xlsx_export():
    """NumPy-backed values used by pandas exports stay numeric in XLSX."""
    rows = [
        [np.int64(2), np.float64(1.25)],
        [np.int64(3), np.float64(2.5)],
        [np.sum(np.array([2, 3], dtype=np.int64)), np.sum([1.25, 2.5])],
    ]

    builder = SpreadsheetBuilder(file_format="xlsx")
    builder.add_sheet(
        sheet_name="Totals",
        columns=[
            SpreadsheetColumn("Units", "int"),
            SpreadsheetColumn("Amount", "float"),
        ],
        rows=rows,
    )
    sheet = load_workbook(BytesIO(builder.build_spreadsheet()), data_only=True)[
        "Totals"
    ]
    values = [
        (sheet.cell(row, 1).value, sheet.cell(row, 2).value) for row in range(2, 5)
    ]

    assert all(
        sheet.cell(row, col).data_type == "n" for row in range(2, 5) for col in (1, 2)
    )
    assert values == [(2, 1.25), (3, 2.5), (5, 3.75)]


def test_native_numeric_values_and_totals_remain_numeric_in_xls_export():
    builder = SpreadsheetBuilder(file_format="xls")
    builder.add_sheet(
        sheet_name="Totals",
        columns=[
            SpreadsheetColumn("Units", "int"),
            SpreadsheetColumn("Amount", "float"),
        ],
        rows=[[2, 1.25], [3, 2.5], [5, 3.75]],
    )
    sheet = xlrd.open_workbook(file_contents=builder.build_spreadsheet()).sheet_by_name(
        "Totals"
    )

    assert all(
        sheet.cell(row, col).ctype == xlrd.XL_CELL_NUMBER
        for row in range(1, 4)
        for col in (0, 1)
    )
    assert [
        (sheet.cell(row, 0).value, sheet.cell(row, 1).value) for row in range(1, 4)
    ] == [(2, 1.25), (3, 2.5), (5, 3.75)]


def test_numpy_numeric_scalars_and_totals_remain_numeric_in_xls_export():
    rows = [
        [np.int64(2), np.float64(1.25)],
        [np.int64(3), np.float64(2.5)],
        [np.sum(np.array([2, 3], dtype=np.int64)), np.sum([1.25, 2.5])],
    ]
    builder = SpreadsheetBuilder(file_format="xls")
    builder.add_sheet(
        sheet_name="Totals",
        columns=[
            SpreadsheetColumn("Units", "int"),
            SpreadsheetColumn("Amount", "float"),
        ],
        rows=rows,
    )
    sheet = xlrd.open_workbook(file_contents=builder.build_spreadsheet()).sheet_by_name(
        "Totals"
    )

    assert all(
        sheet.cell(row, col).ctype == xlrd.XL_CELL_NUMBER
        for row in range(1, 4)
        for col in (0, 1)
    )
    assert [
        (sheet.cell(row, 0).value, sheet.cell(row, 1).value) for row in range(1, 4)
    ] == [(2, 1.25), (3, 2.5), (5, 3.75)]
